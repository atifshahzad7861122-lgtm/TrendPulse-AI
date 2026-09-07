import logging
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

def export_search_results_to_excel(products, output_path):
    """
    Exports a list of product search results to a formatted Excel spreadsheet.
    """
    if not products:
        logger.warning("No product data to export.")
        return False
        
    try:
        df = pd.DataFrame(products)
        writer = pd.ExcelWriter(output_path, engine='openpyxl')
        df.to_excel(writer, index=False, sheet_name='Daraz Products')
        
        workbook = writer.book
        worksheet = writer.sheets['Daraz Products']
        
        header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
        cell_font = Font(name='Segoe UI', size=10)
        
        left_align = Alignment(horizontal='left', vertical='center')
        center_align = Alignment(horizontal='center', vertical='center')
        
        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )
        
        for col_idx in range(1, len(df.columns) + 1):
            cell = worksheet.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border
            
        for row_idx in range(2, worksheet.max_row + 1):
            worksheet.views.sheetView[0].showGridLines = True
            for col_idx in range(1, len(df.columns) + 1):
                cell = worksheet.cell(row=row_idx, column=col_idx)
                cell.font = cell_font
                cell.border = thin_border
                
                col_name = df.columns[col_idx - 1]
                if col_name in ['Rating', 'Reviews', 'Location', 'Discount']:
                    cell.alignment = center_align
                else:
                    cell.alignment = left_align
                    
        for col in worksheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or '')
                if len(val) > 50 and col_letter in ['I', 'J']:
                    val = val[:40] + "..."
                if len(val) > max_len:
                    max_len = len(val)
            worksheet.column_dimensions[col_letter].width = max(max_len + 3, 12)
            
        writer.close()
        logger.info(f"Successfully exported search results to {output_path}")
        return True
    except Exception as e:
        logger.error(f"Failed to export search results: {e}")
        return False


def export_to_excel(product_data: dict, specs: dict, reviews: list, output_filepath: str = "daraz_scraped_data.xlsx"):
    """
    Exports scraped Daraz product details, specs, and reviews into a formatted Excel workbook.
    """
    try:
        wb = openpyxl.Workbook()
        
        # Styles
        navy_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        # -------------------------------------------------------------
        # Sheet 1: Overview
        # -------------------------------------------------------------
        ws_overview = wb.active
        ws_overview.title = "Product Overview"
        ws_overview.views.sheetView[0].showGridLines = True

        ws_overview['A1'] = "Daraz Product Scraping Report"
        ws_overview['A1'].font = Font(name="Segoe UI", size=16, bold=True, color="1E293B")

        overview_rows = [
            ("Attribute", "Extracted Value"),
            ("Product Title", product_data.get("title", "N/A")),
            ("Current Price", product_data.get("price", "N/A")),
            ("Original Price", product_data.get("original_price", "N/A")),
            ("Brand", product_data.get("brand", "N/A")),
            ("Seller Name", product_data.get("seller_name", "N/A")),
            ("Product URL", product_data.get("url", "N/A")),
            ("Total Reviews Scraped", len(reviews))
        ]

        for k, v in product_data.get("seller_metrics", {}).items():
            overview_rows.append((f"Seller Metric: {k}", v))

        for r_idx, (k, v) in enumerate(overview_rows, start=3):
            c1 = ws_overview.cell(row=r_idx, column=1, value=k)
            c2 = ws_overview.cell(row=r_idx, column=2, value=v)
            if r_idx == 3:
                c1.fill, c2.fill = navy_fill, navy_fill
                c1.font, c2.font = header_font, header_font
            else:
                c1.border, c2.border = thin_border, thin_border
                c1.font = Font(name="Segoe UI", size=10, bold=True, color="334155")
                c2.font = Font(name="Segoe UI", size=10, color="0F172A")
                if r_idx % 2 == 0:
                    c1.fill, c2.fill = zebra_fill, zebra_fill

        # -------------------------------------------------------------
        # Sheet 2: Specifications
        # -------------------------------------------------------------
        ws_specs = wb.create_sheet(title="Specifications")
        ws_specs.views.sheetView[0].showGridLines = True
        ws_specs.append(["Specification Attribute", "Specification Value"])
        for col in (1, 2):
            cell = ws_specs.cell(row=1, column=col)
            cell.fill, cell.font = navy_fill, header_font

        for r_idx, (k, v) in enumerate(specs.items(), start=2):
            ws_specs.append([k, v])
            ws_specs.cell(row=r_idx, column=1).border = thin_border
            ws_specs.cell(row=r_idx, column=2).border = thin_border
            if r_idx % 2 == 1:
                ws_specs.cell(row=r_idx, column=1).fill = zebra_fill
                ws_specs.cell(row=r_idx, column=2).fill = zebra_fill

        # -------------------------------------------------------------
        # Sheet 3: Customer Reviews
        # -------------------------------------------------------------
        ws_reviews = wb.create_sheet(title="Customer Reviews")
        ws_reviews.views.sheetView[0].showGridLines = True
        rev_headers = ["Page #", "Reviewer Name", "Date", "Star Rating", "Variation", "Review Text", "Images"]
        ws_reviews.append(rev_headers)

        for col in range(1, len(rev_headers) + 1):
            cell = ws_reviews.cell(row=1, column=col)
            cell.fill, cell.font = navy_fill, header_font
            cell.alignment = Alignment(horizontal="center" if col in [1, 3, 4] else "left")

        for r_idx, rev in enumerate(reviews, start=2):
            row_vals = [
                rev.get("page", 1),
                rev.get("reviewer", "Anonymous"),
                rev.get("date", "N/A"),
                rev.get("rating", 5),
                rev.get("variation", "N/A"),
                rev.get("content", ""),
                rev.get("images", "")
            ]
            ws_reviews.append(row_vals)
            for c_idx in range(1, len(row_vals) + 1):
                cell = ws_reviews.cell(row=r_idx, column=c_idx)
                cell.border = thin_border
                if c_idx in [1, 3, 4]:
                    cell.alignment = Alignment(horizontal="center")
                if r_idx % 2 == 1:
                    cell.fill = zebra_fill

        # Adjust column sizing
        for ws in [ws_overview, ws_specs, ws_reviews]:
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = min(max(max_len + 4, 14), 65)

        wb.save(output_filepath)
        logger.info(f"Successfully exported single product report to {output_filepath}")
        return output_filepath
    except Exception as e:
        logger.error(f"Failed to export single product report to Excel: {e}")
        return False
